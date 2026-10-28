import openpyxl
import json
import os
from collections import Counter

def extract_county_deep_strategies():
    files = {
        'Lake': 'data/stan_strategies/Lake_2025_Sale_Revised_Bid_Totals Claude Analysis.xlsx',
        'DuPage': 'data/stan_strategies/DuPage_2025_Sale_Revised_Bid_Totals- Claude Analysis.xlsx',
        'McHenry': 'data/stan_strategies/McHenry_2025_Sale_Revised_Bid_Totals Claude Analysis.xlsx',
        'Will': 'data/stan_strategies/Will_2025_Sale_Revised_Bid_Totals Claude Analysis.xlsx'
    }

    county_data = {}

    for county, fpath in files.items():
        if not os.path.exists(fpath):
            continue

        wb = openpyxl.load_workbook(fpath, data_only=True)
        
        # 1. Summary
        ws_sum = wb['Summary']
        rule = str(ws_sum.cell(2, 1).value or '')
        
        breakdowns = []
        for r in range(4, ws_sum.max_row + 1):
            c1 = ws_sum.cell(r, 1).value
            c2 = ws_sum.cell(r, 2).value
            c3 = ws_sum.cell(r, 3).value
            if c1 is not None:
                cnt = int(c2) if c2 is not None and str(c2).isdigit() else str(c2 or '')
                amt = float(c3) if c3 is not None and isinstance(c3, (int, float)) else str(c3 or '')
                breakdowns.append({
                    "label": str(c1).strip(),
                    "count": cnt,
                    "amount": amt
                })

        # 2. Non-Comm Sheet Analysis
        ws_nc = wb['Revised Non-Commercial']
        nc_headers = [str(ws_nc.cell(1, c).value or '').strip() for c in range(1, ws_nc.max_column + 1)]
        
        # Find Bid col
        bid_idx = None
        for i, h in enumerate(nc_headers):
            if h in ['Bid', 'TempBid']:
                bid_idx = i + 1
                break

        # Find Comment col
        comm_idx = None
        for i, h in enumerate(nc_headers):
            if h in ['Comment', '2024Comment', 'Message']:
                comm_idx = i + 1
                break

        bids = []
        comments = []
        for r in range(2, ws_nc.max_row + 1):
            if bid_idx:
                bv = ws_nc.cell(r, bid_idx).value
                if bv is not None:
                    bids.append(str(bv).strip().replace('%', ''))
            if comm_idx:
                cv = ws_nc.cell(r, comm_idx).value
                if cv is not None and str(cv).strip() not in ['None', '', 'False', 'True', '0']:
                    comments.append(str(cv).strip())

        bid_counts = Counter(bids)
        comm_counts = Counter(comments)

        county_data[county] = {
            "county": county,
            "rule": rule,
            "breakdowns": breakdowns,
            "bid_distribution": dict(bid_counts.most_common(5)),
            "top_comments": dict(comm_counts.most_common(5)),
            "non_comm_count": ws_nc.max_row - 1,
            "comm_count": wb['Commercial'].max_row - 1,
            "removed_count": wb['Removed - Full Inst Under 3K'].max_row - 1
        }

    with open('stan_county_deep_strategies.json', 'w') as f:
        json.dump(county_data, f, indent=2)

    print("✅ Successfully generated stan_county_deep_strategies.json!")

if __name__ == '__main__':
    extract_county_deep_strategies()
