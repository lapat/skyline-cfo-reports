import openpyxl
import json
import os

def parse_stan_rollup():
    rollup_path = 'data/stan_strategies/4_County_2025_Sale_Revised_Rollup Claude Analysis.xlsx'
    if not os.path.exists(rollup_path):
        print(f"Error: {rollup_path} not found.")
        return

    wb = openpyxl.load_workbook(rollup_path, data_only=True)
    ws = wb['4 County Roll-Up']

    # Row 6: Lake, Row 7: DuPage, Row 8: McHenry, Row 9: Will, Row 10: Total
    # Columns:
    # A: County
    # B: Non-Comm Count, C: Non-Comm Amount
    # D: Comm Count, E: Comm Amount
    # F: Revised Total Count, G: Revised Total Amount
    # H: Removed Count, I: Removed Amount
    # J: Orig Count, K: Orig Amount

    counties = []
    for r in range(6, 11):
        c_name = ws.cell(r, 1).value
        non_comm_count = int(ws.cell(r, 2).value or 0)
        non_comm_amt = float(ws.cell(r, 3).value or 0.0)
        comm_count = int(ws.cell(r, 4).value or 0)
        comm_amt = float(ws.cell(r, 5).value or 0.0)
        rev_count = int(ws.cell(r, 6).value or 0)
        rev_amt = float(ws.cell(r, 7).value or 0.0)
        rem_count = int(ws.cell(r, 8).value or 0)
        rem_amt = float(ws.cell(r, 9).value or 0.0)
        orig_count = int(ws.cell(r, 10).value or 0)
        orig_amt = float(ws.cell(r, 11).value or 0.0)

        entry = {
            "county": c_name,
            "non_commercial": {"count": non_comm_count, "amount": non_comm_amt},
            "commercial": {"count": comm_count, "amount": comm_amt},
            "revised_total": {"count": rev_count, "amount": rev_amt},
            "removed_under_3k": {"count": rem_count, "amount": rem_amt},
            "original_total": {"count": orig_count, "amount": orig_amt}
        }
        counties.append(entry)

    out_data = {
        "title": "2025 Tax Sale – Revised Bid Totals, 4 County Roll-Up",
        "description": "Stan's production bidding models. Removed = certs under $3,000 with both installments delinquent. Commercial includes industrial.",
        "rollup": counties
    }

    out_file = 'stan_bidding_models.json'
    with open(out_file, 'w') as f:
        json.dump(out_data, f, indent=2)
    print(f"✅ Successfully wrote Stan's bidding models to {out_file}")

if __name__ == '__main__':
    parse_stan_rollup()
