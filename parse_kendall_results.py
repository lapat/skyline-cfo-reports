import openpyxl
import json
import os

def parse_kendall():
    fpath = 'data/kendall_auction_results/Kendall_Results All Buyers 2026 Sale 2025 Taxes.xlsx'
    if not os.path.exists(fpath):
        print(f"Error: {fpath} not found")
        return

    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws_tot = wb['Totals by Bidder']

    buyers = []
    for r in range(4, ws_tot.max_row + 1):
        c_num = ws_tot.cell(r, 1).value
        c_name = ws_tot.cell(r, 2).value
        c_certs = ws_tot.cell(r, 3).value
        c_amt = ws_tot.cell(r, 4).value
        c_pct = ws_tot.cell(r, 5).value
        c_wtd = ws_tot.cell(r, 6).value
        c_min = ws_tot.cell(r, 7).value
        c_max = ws_tot.cell(r, 8).value

        if c_name and c_name != 'None':
            # Check if total
            is_tot = 'total' in str(c_name).lower() or 'total' in str(c_num).lower()
            try:
                amt_val = float(c_amt or 0)
                certs_val = int(c_certs or 0)
                pct_val = float(c_pct or 0) * 100 if float(c_pct or 0) <= 1.0 else float(c_pct or 0)
                wtd_val = float(c_wtd or 0)
            except:
                continue

            buyers.append({
                "buyer_num": str(c_num or '').strip(),
                "buyer_name": str(c_name).strip(),
                "certs_won": certs_val,
                "amount_won": round(amt_val, 2),
                "share_pct": round(pct_val, 1),
                "wtd_avg_bid": round(wtd_val, 2),
                "min_bid": str(c_min or ''),
                "max_bid": str(c_max or ''),
                "is_total": is_tot
            })

    # Ingest Jennings bid summary from Stan's email
    jennings_bid_summary = {
        "bidder_name": "Interstate Funding Corp (Jennings)",
        "buyer_code": "042",
        "items_bid": 179,
        "total_amount_bid": 1495880.80,
        "items_won": 0,
        "amount_won": 0.0,
        "bid_distribution": [
            {"rate": "6%", "pins": 35, "amount": 248243.43},
            {"rate": "7%", "pins": 64, "amount": 502280.67},
            {"rate": "8%", "pins": 76, "amount": 564806.91},
            {"rate": "9%", "pins": 4, "amount": 180549.79}
        ],
        "outcome": "ZERO WINS (0.0% Clearance)",
        "lesson": "Bidding 6%-9% resulted in 0 wins because Bridge took 50% of the sale at 4.95% avg, and Bass took 13% at 4.12% avg."
    }

    data = {
        "title": "Kendall County 2026 Tax Sale Results (2025 Taxes)",
        "sale_date": "2026-10-26",
        "total_certificates_sold": 417,
        "total_amount_sold": 2291638.01,
        "market_wtd_avg_bid": 4.88,
        "jennings_bids": jennings_bid_summary,
        "buyers": buyers,
        "jay_directive": "We are going to bid 5 moving forward so that will be our rate and I will send over the bid ladder so once we have it all approved, we just put in 5."
    }

    out_file = 'data/kendall_auction_results/kendall_results_2026.json'
    with open(out_file, 'w') as f:
        json.dump(data, f, indent=2)

    print(f"✅ Successfully wrote Kendall results to {out_file}")
    print(f"   Total Market Cleared: {data['total_certificates_sold']} certs (${data['total_amount_sold']:,.2f}) @ {data['market_wtd_avg_bid']}% avg")
    print(f"   Jennings Result: Bid 179 items (${jennings_bid_summary['total_amount_bid']:,.2f}) @ 6-9% -> Won 0 items ($0.00)")

if __name__ == '__main__':
    parse_kendall()
