import openpyxl
import json
import os
import glob

def parse_all_stan_data():
    files = {
        'Lake': 'data/stan_strategies/Lake_2025_Sale_Revised_Bid_Totals Claude Analysis.xlsx',
        'DuPage': 'data/stan_strategies/DuPage_2025_Sale_Revised_Bid_Totals- Claude Analysis.xlsx',
        'McHenry': 'data/stan_strategies/McHenry_2025_Sale_Revised_Bid_Totals Claude Analysis.xlsx',
        'Will': 'data/stan_strategies/Will_2025_Sale_Revised_Bid_Totals Claude Analysis.xlsx'
    }

    commercial_targets = []
    category_summaries = {}
    redline_parcels = []

    for county, fpath in files.items():
        if not os.path.exists(fpath):
            print(f"Missing {fpath}")
            continue

        wb = openpyxl.load_workbook(fpath, data_only=True)
        
        # 1. Parse Summary Tab
        ws_sum = wb['Summary']
        sum_rows = []
        for r in range(4, ws_sum.max_row + 1):
            c1 = ws_sum.cell(r, 1).value
            c2 = ws_sum.cell(r, 2).value
            c3 = ws_sum.cell(r, 3).value
            if c1 is not None:
                sum_rows.append({
                    "category": str(c1).strip(),
                    "count": int(c2) if c2 is not None and str(c2).isdigit() else str(c2 or ''),
                    "amount": float(c3) if c3 is not None and isinstance(c3, (int, float)) else str(c3 or '')
                })
        category_summaries[county] = sum_rows

        # 2. Parse Commercial Tab
        ws_comm = wb['Commercial']
        headers = [str(ws_comm.cell(1, c).value or '').strip() for c in range(1, ws_comm.max_column + 1)]
        
        for r in range(2, ws_comm.max_row + 1):
            row = {headers[c-1]: ws_comm.cell(r, c).value for c in range(1, ws_comm.max_column + 1) if headers[c-1]}
            
            # Normalize fields
            pin = str(row.get('Parcel Number') or row.get('Parcel') or row.get('PINNoDashes') or '').strip()
            taxpayer = str(row.get('Taxpayer Name') or row.get('ASSESSEE NAME') or row.get('Name') or '').strip()
            address = str(row.get('Property Address 1') or row.get('PROPERTY ADDRESS') or row.get('SiteAddress') or '').strip()
            city = str(row.get('Property City') or row.get('SITECITY') or '').strip()
            
            # Tax Due
            tax_due = (row.get('Judgment Amount') or row.get('SaleAmount') or row.get('Total-amount-owed') or 0.0)
            if isinstance(tax_due, str):
                tax_due = float(tax_due.replace(',', '').replace('$', '').strip() or 0.0)
            else:
                tax_due = float(tax_due or 0.0)

            # Bid
            bid = str(row.get('Bid') or row.get('TempBid') or '3%').strip()
            if bid.isdigit():
                bid = f"{bid}%"

            # Comment
            comment = str(row.get('Comment') or row.get('2024Comment') or row.get('Message') or '').strip()
            if comment in ['None', 'False', 'True', '0']:
                comment = ''

            # Asmt Bldg & Land
            asmt_bldg = float(row.get('Asmt Bldg') or row.get('Assessed Bldg') or row.get('NONFARMBLD') or row.get('buildingAssd') or 0.0)
            asmt_land = float(row.get('Asmt Land') or row.get('Assessed Land') or row.get('NONFARMLND') or row.get('landAssd') or 0.0)

            # NoMort
            no_mort = row.get('NoMort') in [True, 'True', 'Y', 1] or (row.get('MortAmt') in [0, 0.0, None, 'None', ''])

            # HalfInstallment
            half_inst = row.get('1Inst') in [True, 'True', 'Y', 1] or row.get('HalfInstallment') in [True, 'True', 'Y', 1] or row.get('OneInstllmt') in [True, 'True', 'Y', 1]

            if pin and tax_due > 0:
                commercial_targets.append({
                    "county": county,
                    "pin": pin,
                    "taxpayer": taxpayer,
                    "address": address,
                    "city": city,
                    "tax_due": tax_due,
                    "bid": bid,
                    "comment": comment,
                    "asmt_bldg": asmt_bldg,
                    "asmt_land": asmt_land,
                    "bldg_land_ratio": round(asmt_bldg / asmt_land, 1) if asmt_land > 0 else 0.0,
                    "no_mortgage": no_mort,
                    "is_half_installment": half_inst
                })

    # Sort commercial targets by tax_due descending
    commercial_targets.sort(key=lambda x: x['tax_due'], reverse=True)

    print(f"✅ Extracted {len(commercial_targets)} commercial target parcels across 4 counties.")
    tot_comm_capital = sum(c['tax_due'] for c in commercial_targets)
    print(f"   Total Commercial Capital: ${tot_comm_capital:,.2f}")

    # Write commercial targets JSON
    with open('stan_commercial_targets.json', 'w') as f:
        json.dump(commercial_targets, f, indent=2)

    # Write category summaries JSON
    with open('stan_category_summaries.json', 'w') as f:
        json.dump(category_summaries, f, indent=2)

    print("✅ Wrote stan_commercial_targets.json and stan_category_summaries.json successfully!")

if __name__ == '__main__':
    parse_all_stan_data()
